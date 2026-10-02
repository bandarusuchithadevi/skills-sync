import io
import zipfile
import docx

def is_docx(file_bytes: bytes) -> bool:
    """Check if raw bytes represent a valid DOCX zip structure."""
    if not file_bytes or not file_bytes.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            return "word/document.xml" in zf.namelist()
    except Exception:
        return False

def extract_text_from_docx(file_bytes: bytes) -> tuple[str, int]:
    """Extract text and estimated page count from raw DOCX bytes."""
    if not file_bytes or len(file_bytes.strip()) == 0:
        raise ValueError("File is empty.")

    if not is_docx(file_bytes):
        raise ValueError("File is not a valid DOCX document.")

    try:
        doc_file = io.BytesIO(file_bytes)
        document = docx.Document(doc_file)
        full_text = []
        for p in document.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())

        for table in document.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))

        text = "\n".join(full_text).strip()
        if not text:
            raise ValueError("Could not extract readable text from DOCX file. The document appears to be empty.")

        word_count = len(text.split())
        estimated_pages = max(1, (word_count + 299) // 300)
        return text, estimated_pages
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")
